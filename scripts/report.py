# -*- coding: utf-8 -*-
"""이벤트 스터디 pptx 보고서 — 슬라이드 수 = 1 + 2N (CLAUDE.md 6.2).

  1장        조건과 사례 목록
  사례당 2장  (a) 지표 해부  (b) 시장 반응

실행: py -3 scripts/report.py
산출: _out/report.pptx  (+ 검증용 _out/render/*.PNG)
"""
import sys
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import glob
import json
import os
import shutil

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pptx_kit as K

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "_out")
ANAT = os.path.join(ROOT, "data", "anatomy")

TODAY = "2026-09-04"
TITLE = "근원 CPI 상회 · 고용 서프라이즈 · 제조업 확장이 동시 성립한 국면"

NAMES = {"^GSPC": "S&P500", "^IXIC": "나스닥", "^SOX": "SOX", "^RUT": "러셀2000",
         "CYC": "경기민감", "DEF": "방어"}
IDX = ["^GSPC", "^IXIC", "^SOX", "^RUT"]          # 1장 메인 표는 지수 4종만 (6.2)
ALL6 = IDX + ["CYC", "DEF"]                        # 사례 장은 바스켓 포함
HD = ["0", "1", "5", "20", "60"]
HLAB = ["당일", "1일", "5일", "20일", "60일"]

# 열 폭 (inch) — 합이 본문 폭에 정확히 맞는다
COND_COLW = [0.85, 2.00, 1.50, 2.50, 5.483]
MAIN_COLW = [1.45, 1.05, 1.05, 0.80, 1.35, 1.05, 1.05, 1.05, 1.05, 1.05, 1.383]
ANAT_X = [0.50, 4.85, 8.30]
ANAT_COLW = [[1.85, 0.75, 0.70, 0.85],
             [1.40, 0.875, 0.875],
             [1.25, 0.70, 0.70, 0.65, 1.233]]
RET_COLW = [1.55, 0.72, 0.70, 0.70, 0.74, 0.74, 0.85]
LEFT_X, LEFT_W = 0.50, 6.00
RIGHT_X, RIGHT_W = 6.73, 6.10

ROWH, HDRH = 0.29, 0.32


def pct(v):
    return "-" if v is None else "%+.2f" % v


def _t(x):
    """anatomy 의 문장은 {text, from} 객체다 (근거 대조용). 옛 문자열도 받는다."""
    return x if isinstance(x, str) else x["text"]


# ── 1장 ──────────────────────────────────────────────────────────────────
def slide_overview(prs, res, anat):
    sl = K.blank(prs)
    K.band(sl, "%s  |  이벤트 스터디" % TODAY, TITLE,
           "대상 기간 %s ~ %s  ·  수익률 기산일은 세 조건 중 가장 늦게 발표되는 근원 CPI 발표일"
           % tuple(res["data_span"]["cpi"]))

    y = K.h2(sl, K.M, 1.14, K.BODY_W, "조건 정의")
    tbl, y = K.table(sl, K.M, y, COND_COLW, 4, 0.27, HDRH)
    K.header_row(tbl, ["조건", "지표", "variant", "임계값", "역할"])
    rows = [["조건 1", "근원 CPI", "core MoM", "surprise ≥ +0.20%p", "수익률 기산일 — 이 지표의 발표일이 t=0"],
            ["조건 2", "비농업고용 (NFP)", "headline", "surprise ≥ +100K", "판정만. 수익률 계산에는 쓰지 않음"],
            ["조건 3", "ISM 제조업 PMI", "headline", "actual ≥ 50.0", "판정만. 서프라이즈 아닌 레벨 조건"]]
    for i, r in enumerate(rows):
        K.body_row(tbl, i + 1, r, ["c", "l", "l", "l", "l"],
                   last=(i == len(rows) - 1), sign_from=99)

    K.text(sl, K.M, y + 0.08, K.BODY_W, 0.6, [
        "분석 대상 — S&P500 · 나스닥 · SOX · 러셀2000, 경기민감 바스켓(XLI · XLY · XLF · XLB 동일가중) · 방어 바스켓(XLP · XLU · XLV 동일가중)",
        "관찰 구간 — 당일(전일 종가 → 당일 종가) · 1 · 5 · 20 · 60거래일. CPI는 08:30 ET 장 시작 전 발표라 당일 종가가 첫 반응임",
        "결합 방식 — 동일 참조월로 세 지표를 정렬. 근원 CPI가 셋 중 가장 늦게 발표되므로 기산일 시점에 NFP · ISM은 이미 공개 → 룩어헤드 없음. 컨센서스 출처는 Investing.com",
    ], size=K.F_FOOT + 0.5, color=K.GRAY, space=1)

    y = K.h2(sl, K.M, 3.26, K.BODY_W, "조건 충족일")
    f = res["funnel"]
    K.text(sl, K.M, y, K.BODY_W, 0.22,
           "퍼널 — 전체 %d개월 → 조건 1 통과 %d건 → 조건 2까지 %d건 → 조건 3까지 %d건. "
           "조건 1 통과 %d건은 ISM이 모두 50을 넘어 조건 3은 제약으로 작동하지 않았음 — 표본을 결정한 것은 조건 2임"
           % (f["all"], f["A"], f["B"], f["C"], f["A"]), size=K.F_TBL, space=0)

    cases = res["cases"]
    n = len(cases) * len(IDX)
    tbl, y = K.table(sl, K.M, y + 0.28, MAIN_COLW, n + 1, 0.27, HDRH)
    K.header_row(tbl, ["기산일", "CPI 서프", "NFP 서프", "ISM", "대상"] + HLAB + ["60일 MDD"])

    for ci, c in enumerate(cases):
        r0 = 1 + ci * len(IDX)
        last_blk = (ci == len(cases) - 1)
        for j in range(4):                                   # 병합 열
            tbl.cell(r0, j).merge(tbl.cell(r0 + len(IDX) - 1, j))
        edge = (1.0, K.NAVY) if last_blk else (0.5, K.RULE)
        for j, v in enumerate([c["date"], "%+.2f" % c["cpi_surp"],
                               "%+.0f" % c["nfp_surp"], "%.1f" % c["ism"]]):
            K.set_cell(tbl, r0, j, v, bold=True, align="c" if j else "l",
                       auto_sign=(j in (1, 2)), bottom=edge)
        for k, t in enumerate(IDX):
            r = r0 + k
            last = last_blk and k == len(IDX) - 1
            bot = (1.0, K.NAVY) if last else (0.5, K.RULE)
            K.set_cell(tbl, r, 4, NAMES[t], align="l", bottom=bot)
            for h, col in zip(HD, range(5, 10)):
                K.set_cell(tbl, r, col, pct(c["ret"][t][h]), auto_sign=True, bottom=bot)
            K.set_cell(tbl, r, 10, pct(c["mdd60"][t]), auto_sign=True, bottom=bot)

    bl = res["baseline"]
    base = " · ".join("%s %s/%s/%s/%s" % (
        NAMES[t], pct(bl[t]["1"]["median"]), pct(bl[t]["5"]["median"]),
        pct(bl[t]["20"]["median"]), pct(bl[t]["60"]["median"])) for t in IDX)
    K.text(sl, K.M, y + 0.07, K.BODY_W, 0.4, [
        "단위 — CPI 서프 %p, NFP 서프 K, 수익률 · MDD %. 컨센서스 Investing.com, 가격 Yahoo Finance. 개별종목이 아닌 지수 · ETF만 사용해 생존편향은 회피",
        "무조건부 베이스라인(%s ~ %s) 중앙값 1/5/20/60일 — %s"
        % (res["data_span"]["base"][0], res["data_span"]["base"][1], base),
    ], size=K.F_FOOT, color=K.GRAY, space=1)

    K.warn(sl, K.M, 6.72, K.BODY_W, 0.46, [
        "n = %d — 통계적 유의성 없음. 사례 참고용. 중앙값 · 상승확률 등 요약통계는 산출하지 않고 %d건의 원본 경로만 제시함"
        % (f["C"], f["C"]),
        "두 사례 간격 %d일로 60거래일 관찰 구간이 겹치지 않아 자기상관 문제는 없으나, 표본이 모두 2020~2021년 코로나 재개방기에 속해 다른 레짐으로의 일반화 근거는 없음"
        % res["gaps_days"][0],
    ])


# ── 사례 (a) 지표 해부 ───────────────────────────────────────────────────
def slide_anatomy(prs, c, a, i, ntot):
    sl = K.blank(prs)
    K.band(sl, "%s  |  이벤트 스터디 · 사례 %d / %d" % (TODAY, i + 1, ntot),
           "[사례 %d] %s — 지표 해부" % (i + 1, c["date"]),
           "참조월 {}  ·  근원 CPI {:.1f}% (컨센 {:.1f}%)  ·  NFP {:+,.0f}K (컨센 {:+,.0f}K)  ·  ISM {:.1f}".format(
               c["ref"], c["cpi_act"], c["cpi_fc"], c["nfp_act"], c["nfp_fc"], c["ism"]))

    for j, blk in enumerate(a["anatomy"]):
        x, colw = ANAT_X[j], ANAT_COLW[j]
        w = sum(colw)
        y = K.h2(sl, x, 1.18, w, "%s · %s" % (blk["cond"], blk["name"]))
        K.text(sl, x, y - 0.02, w, 0.20, blk["headline"], size=K.F_TBL, color=K.GRAY, space=0)
        tbl, ty = K.table(sl, x, 1.80, colw, len(blk["rows"]) + 1, ROWH, HDRH)
        K.header_row(tbl, blk["cols"])
        for r, row in enumerate(blk["rows"]):
            last = (r == len(blk["rows"]) - 1)
            gray = "잔차" in row[0]
            K.body_row(tbl, r + 1, row, blk["align"], last=last, bold=last)
            if gray:
                K.set_cell(tbl, r + 1, 0, row[0], align="l", color=K.GRAY,
                           bottom=(0.5, K.RULE))
        K.text(sl, x, 5.12, w, 0.34, blk["src"], size=K.F_FOOT, color=K.GRAY, space=0)

    lines = []
    for blk in a["anatomy"]:
        for k, nt in enumerate(blk["note"]):
            lines.append((("%s · %s   " % (blk["cond"], blk["name"].split(" (")[0])
                           if k == 0 else "        ") + _t(nt),
                          {"bold": False}))
    K.text(sl, K.M, 5.62, K.BODY_W, 1.5, lines, size=K.F_TBL, space=2)


# ── 사례 (b) 시장 반응 ───────────────────────────────────────────────────
def slide_market(prs, c, a, i, ntot):
    sl = K.blank(prs)
    K.band(sl, "%s  |  이벤트 스터디 · 사례 %d / %d" % (TODAY, i + 1, ntot),
           "[사례 %d] %s — 시장 반응" % (i + 1, c["date"]),
           "기산일 당일 종가 기준 · 경기민감/방어는 섹터 ETF 동일가중 바스켓")

    y = K.h2(sl, LEFT_X, 1.18, LEFT_W, "구간 수익률")
    tbl, y = K.table(sl, LEFT_X, 1.60, RET_COLW, len(ALL6) + 1, ROWH, HDRH)
    K.header_row(tbl, ["대상"] + HLAB + ["60일 MDD"])
    for k, t in enumerate(ALL6):
        vals = [NAMES[t]] + [pct(c["ret"][t][h]) for h in HD] + \
               [pct(c["mdd60"][t]) if t in c["mdd60"] else "-"]
        K.body_row(tbl, k + 1, vals, ["l"] + ["r"] * 6, last=(k == len(ALL6) - 1))

    y = K.h2(sl, LEFT_X, 3.68, LEFT_W, "당일 시장 반응")
    K.text(sl, LEFT_X, y, LEFT_W, 0.8, [_t(b) for b in a["day"]],
           size=K.F_TBL, space=2, bullet=True)

    y = K.h2(sl, LEFT_X, 4.75, LEFT_W, "이후 경로")
    K.text(sl, LEFT_X, y, LEFT_W, 1.6,
           [(_t(b), {}) for b in a["path"]] +
           [(_t(a["read"]), {"mark": "▷ ", "color": K.NAVY, "bold": True})],
           size=K.F_TBL, space=2, bullet=True)

    fy, box = 1.14, (7.18 - 1.14 - 2 * (0.22 + 0.20)) / 2.0
    for n, (fn, cap, src) in enumerate([
            ("case%d_path.png" % (i + 1), "[그림 %d-1] 이벤트 전후 누적 경로" % (i + 1),
             "자료: Yahoo Finance, 자체 계산. 발표일 종가 = 0%로 정규화, −20 ~ +60거래일"),
            ("case%d_bar.png" % (i + 1), "[그림 %d-2] 대상별 구간 수익률" % (i + 1),
             "자료: Investing.com, Yahoo Finance, 자체 계산")]):
        K.text(sl, RIGHT_X, fy, RIGHT_W, 0.20, cap, size=K.F_TBL, bold=True, space=0)
        fy = K.picture(sl, os.path.join(OUT, fn), RIGHT_X, fy + 0.22, RIGHT_W, box) + 0.03
        K.text(sl, RIGHT_X, fy, RIGHT_W, 0.18, src, size=K.F_FOOT, color=K.GRAY, space=0)
        fy += 0.24


# ── 검증 ─────────────────────────────────────────────────────────────────
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def verify(path, ncase):
    prs = Presentation(path)
    want = 1 + 2 * ncase
    small, outside, no_ea = [], [], 0
    for si, sl in enumerate(prs.slides, 1):
        for shp in sl.shapes:
            if shp.left is None or shp.width is None:
                continue
            l, t = shp.left / 914400.0, shp.top / 914400.0
            r, b = l + shp.width / 914400.0, t + shp.height / 914400.0
            if l < -0.01 or t < -0.01 or r > K.W + 0.01 or b > K.H + 0.01:
                outside.append("s%d %s L%.2f T%.2f R%.2f B%.2f"
                               % (si, shp.shape_type, l, t, r, b))
            for run, in_tbl in _runs(shp):
                if run.font.size is None:
                    continue
                lim = 9 if in_tbl else 8
                if run.font.size.pt < lim:
                    small.append("s%d %.1fpt '%s'" % (si, run.font.size.pt, run.text[:16]))
                rPr = run._r.find(_A + "rPr")
                if rPr is None or rPr.find(_A + "ea") is None:
                    no_ea += 1
    checks = {
        "슬라이드 수 = 1 + 2N (%d)" % want: len(prs.slides) == want,
        "표 run ≥ 9pt · 본문 run ≥ 8pt": not small,
        "모든 도형이 캔버스 안": not outside,
        "모든 run에 ea typeface": no_ea == 0,
    }
    for k, v in checks.items():
        print(("PASS  " if v else "FAIL  ") + k)
    for m in small[:6] + outside[:6]:
        print("      -", m)
    if no_ea:
        print("      - ea 누락 run %d개" % no_ea)
    return all(checks.values())


def _runs(shp):
    if shp.has_text_frame:
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                yield r, False
    if getattr(shp, "has_table", False) and shp.has_table:
        for row in shp.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        yield r, True


def render_png(path):
    """조판 깨짐은 텍스트 검사로 못 잡는다 → PowerPoint 로 PNG 내보내 눈으로 본다."""
    d = os.path.join(OUT, "render")
    shutil.rmtree(d, ignore_errors=True)
    try:
        import win32com.client
        app = win32com.client.Dispatch("PowerPoint.Application")
        pres = app.Presentations.Open(path, WithWindow=False)
        pres.Export(d, "PNG")
        pres.Close()
        app.Quit()
        return sorted(glob.glob(os.path.join(d, "*")))
    except Exception as e:
        print("렌더 실패 (수동 확인 필요):", e)
        return []


def main():
    res = json.load(open(os.path.join(OUT, "results.json"), encoding="utf-8"))
    cases = res["cases"]
    anat = {c["date"]: json.load(open(os.path.join(ANAT, "%s.json" % c["date"]),
                                     encoding="utf-8")) for c in cases}

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(K.W), Inches(K.H)
    slide_overview(prs, res, anat)
    for i, c in enumerate(cases):
        slide_anatomy(prs, c, anat[c["date"]], i, len(cases))
        slide_market(prs, c, anat[c["date"]], i, len(cases))

    path = os.path.join(OUT, "report.pptx")
    prs.save(path)

    ok = verify(path, len(cases))
    pngs = render_png(path)
    print("PPTX:", path, "| slides:", 1 + 2 * len(cases),
          "| render:", len(pngs), "장")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
