# -*- coding: utf-8 -*-
"""pptx 조판 헬퍼 — CLAUDE.md 6.2 디자인 규격을 코드로 고정한 것.

여기 없는 색 · 폰트 크기는 쓰지 않는다. 모든 표 셀은 set_cell() 을 거친다.
"""
import re

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

# ── 색 (6.2) ─────────────────────────────────────────────────────────────
INK, GRAY, NAVY = "1A1A1A", "6B7280", "1F4E79"
HDRBG, WARNBG, RULE = "EEF2F7", "FDF3F2", "DDDDDD"
UP, DOWN = "C0392B", "1F6FB2"          # 국내 관행: 상승 적색 / 하락 청색

# 차트 계열색 — 남색 · 적색 · 회색의 명도 단계 (무지개 금지)
SERIES_COLOR = {"^GSPC": "#1F4E79", "^IXIC": "#3D7AB8", "^SOX": "#C0392B",
                "^RUT": "#E08A82", "CYC": "#6B7280", "DEF": "#B0B6BF"}

# ── 폰트 크기 (6.2) ──────────────────────────────────────────────────────
F_TITLE, F_H2, F_BODY, F_TBL, F_FOOT = 20, 13, 11, 10, 8

# ── 캔버스 (6.2) ─────────────────────────────────────────────────────────
W, H = 13.333, 7.5
M = 0.5                                 # 좌우 여백
BODY_W = W - 2 * M                      # 12.333
BODY_TOP, BODY_BOT = 1.05, 7.15

FONT = "맑은 고딕"
_NUM = re.compile(r"^[+\-−]?[\d,]+(\.\d+)?(%p|%|K|pt)?$")


def _ea(run):
    """한글 글리프용 동아시아 typeface 지정. latin 만으로는 대체폰트가 붙는다."""
    rPr = run._r.get_or_add_rPr()
    for e in rPr.findall(qn("a:ea")):
        rPr.remove(e)
    rPr.append(parse_xml('<a:ea %s typeface="%s"/>' % (nsdecls("a"), FONT)))


def style_run(run, size, bold=False, color=INK):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    _ea(run)
    return run


def sign_color(text, default=INK):
    """수치 문자열의 부호로 글자색을 정한다 (배경 음영 금지)."""
    t = text.strip()
    if _NUM.match(t):
        if t.startswith("+"):
            return UP
        if t.startswith("-") or t.startswith("−"):
            return DOWN
    return default


def minus(text):
    """숫자 셀의 ASCII 하이픈을 U+2212 로 교체. 날짜 문자열은 건드리지 않는다."""
    t = text.strip()
    return "−" + t[1:] if (_NUM.match(t) and t.startswith("-")) else text


def png_size(path):
    """PNG IHDR 에서 픽셀 크기를 읽는다 (PIL 불필요)."""
    with open(path, "rb") as f:
        head = f.read(24)
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def picture(sl, path, x, y, w_max, h_max):
    """종횡비를 지킨 채 (w_max, h_max) 상자 안에 맞춰 넣는다. 하단 y 반환."""
    pw, ph = png_size(path)
    scale = min(w_max / pw, h_max / ph)
    w, h = pw * scale, ph * scale
    sl.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return y + h


def _plain(shape):
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _rect(sl, x, y, w, h, color):
    shp = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    return _plain(shp)


# ── 슬라이드 · 텍스트 ────────────────────────────────────────────────────
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def text(sl, x, y, w, h, lines, size=F_BODY, bold=False, color=INK,
         align="l", space=4, bullet=False):
    """lines = 문자열 또는 (문자열, dict) 목록. dict 로 줄별 서식 덮어쓰기."""
    if isinstance(lines, str):
        lines = [lines]
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        ov = {}
        if isinstance(ln, tuple):
            ln, ov = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[ov.get("align", align)]
        p.space_after = Pt(ov.get("space", space))
        p.line_spacing = 1.18
        r = p.add_run()
        mark = ov.get("mark", "· " if ov.get("bullet", bullet) else "")
        r.text = (mark + ln) if ln else ""
        style_run(r, ov.get("size", size), ov.get("bold", bold), ov.get("color", color))
    return box


def band(sl, kicker, title, sub):
    """상단 타이틀 밴드 (높이 0.9in) — 모든 슬라이드 공통."""
    text(sl, M, 0.22, BODY_W, 0.18, kicker, size=F_FOOT, color=GRAY, space=0)
    text(sl, M, 0.40, BODY_W, 0.36, title, size=F_TITLE, bold=True, space=0)
    text(sl, M, 0.79, BODY_W, 0.20, sub, size=F_BODY - 1, color=GRAY, space=0)
    _rect(sl, M, 1.00, Inches(BODY_W), Pt(1.5), NAVY)


def h2(sl, x, y, w, label):
    """소제목 — 남색 세로 막대 + 텍스트. 다음 요소의 y 를 돌려준다."""
    _rect(sl, x, y + 0.02, Pt(3), Inches(0.20), NAVY)
    text(sl, x + 0.10, y, w - 0.10, 0.24, label, size=F_H2, bold=True, color=NAVY, space=0)
    return y + 0.32


def warn(sl, x, y, w, h, lines):
    _rect(sl, x, y, Inches(w), Inches(h), WARNBG)
    _rect(sl, x, y, Pt(3), Inches(h), UP)
    text(sl, x + 0.13, y + 0.08, w - 0.24, h - 0.16, lines, size=F_FOOT + 0.5, space=1)


# ── 표 ───────────────────────────────────────────────────────────────────
_NOSTYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"      # No Style, No Grid


def _kill_default_style(tbl):
    """python-pptx 가 자동으로 붙이는 파란 줄무늬 테이블 스타일을 제거한다."""
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for e in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(e)
    tblPr.append(parse_xml("<a:tableStyleId %s>%s</a:tableStyleId>"
                           % (nsdecls("a"), _NOSTYLE)))


def _edge(tag, v):
    if v is None:
        return parse_xml("<a:%s %s><a:noFill/></a:%s>" % (tag, nsdecls("a"), tag))
    w_pt, color = v
    return parse_xml(
        '<a:%s %s w="%d" cap="flat" cmpd="sng" algn="ctr">'
        "<a:solidFill><a:srgbClr val=\"%s\"/></a:solidFill>"
        '<a:prstDash val="solid"/></a:%s>'
        % (tag, nsdecls("a"), int(w_pt * 12700), color, tag))


def borders(cell, top=None, bottom=None):
    """세로 괘선은 항상 없음 (6.2). 가로선만 (w_pt, color) 로 지정."""
    tcPr = cell._tc.get_or_add_tcPr()
    spec = [("lnL", None), ("lnR", None), ("lnT", top), ("lnB", bottom)]
    for tag, _ in spec:
        for e in tcPr.findall(qn("a:" + tag)):
            tcPr.remove(e)
    for tag, v in reversed(spec):
        tcPr.insert(0, _edge(tag, v))


def set_cell(tbl, r, c, txt, size=F_TBL, bold=False, align="r",
             color=INK, fill=None, auto_sign=False, top=None, bottom=None):
    """모든 셀은 이 함수를 거친다 (6.2)."""
    cell = tbl.cell(r, c)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        cell.fill.background()
    borders(cell, top, bottom)

    txt = "" if txt is None else str(txt)
    if auto_sign:
        color = sign_color(txt, color)
        txt = minus(txt)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                   "r": PP_ALIGN.RIGHT}[align]
    style_run(p.add_run(), size, bold, color).text = txt
    return cell


def table(sl, x, y, colw, nrows, rowh=0.26, hdr_h=0.30):
    """빈 표 생성 + 기본 스타일 제거. colw 는 inch 목록. (표, 하단 y) 반환."""
    shp = sl.shapes.add_table(nrows, len(colw), Inches(x), Inches(y),
                              Inches(sum(colw)), Inches(hdr_h + (nrows - 1) * rowh))
    tbl = shp.table
    _kill_default_style(tbl)
    for i, w in enumerate(colw):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(hdr_h)
    for i in range(1, nrows):
        tbl.rows[i].height = Inches(rowh)
    return tbl, y + hdr_h + (nrows - 1) * rowh


def header_row(tbl, labels):
    for c, lab in enumerate(labels):
        set_cell(tbl, 0, c, lab, bold=True, align="c", color=NAVY,
                 fill=HDRBG, bottom=(1.25, NAVY))


def body_row(tbl, r, vals, aligns, last=False, bold=False, fill=None, sign_from=1):
    for c, v in enumerate(vals):
        set_cell(tbl, r, c, v, align=aligns[c], bold=bold, fill=fill,
                 auto_sign=(c >= sign_from),
                 bottom=(1.0, NAVY) if last else (0.5, RULE))
